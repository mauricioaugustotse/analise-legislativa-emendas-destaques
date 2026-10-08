from __future__ import annotations

import subprocess
from pathlib import Path

from docx import Document

from analise_legislativa import extract as module


def test_shallow_text_uses_better_ocr(monkeypatch, tmp_path: Path):
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(b"%PDF dummy")
    monkeypatch.setattr(module, "_fitz_pages", lambda _: ["Assinado", "Carimbo"])
    monkeypatch.setattr(module, "_pypdf_pages", lambda _: ["", ""])
    monkeypatch.setattr(module, "_ocr_pages", lambda _: ["Texto reconhecido " * 20, "Outro texto " * 20])
    result = module.extract_pdf(pdf)
    assert result.method == "tesseract-por"
    assert result.segments[1]["page"] == 2


def test_ocr_reads_utf8_and_honors_tesseract_cmd(monkeypatch, tmp_path: Path):
    import pymupdf

    pdf = tmp_path / "scan.pdf"
    with pymupdf.open() as doc:
        doc.new_page().insert_text((72, 72), "imagem")
        doc.save(str(pdf))
    seen = {}

    def fake_run(args, **kwargs):
        seen.update(command=args[0], encoding=kwargs.get("encoding"))
        return subprocess.CompletedProcess(args, 0, stdout="Votação em separado", stderr="")

    monkeypatch.setenv("TESSERACT_CMD", r"C:\Program Files\Tesseract-OCR\tesseract.exe")
    monkeypatch.setattr(module.subprocess, "run", fake_run)
    assert module._ocr_pages(pdf) == ["Votação em separado"]
    assert seen == {"command": r"C:\Program Files\Tesseract-OCR\tesseract.exe", "encoding": "utf-8"}


def test_docx_keeps_table_in_body_order(tmp_path: Path):
    path = tmp_path / "parecer.docx"
    doc = Document()
    doc.add_paragraph("Antes da tabela")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Art. 1º"
    table.cell(0, 1).text = "Regra sintética"
    doc.add_paragraph("Depois da tabela")
    doc.save(path)
    text = module.extract_docx(path).text
    assert text.index("Antes") < text.index("Regra sintética") < text.index("Depois")


def test_authentication_stamp_is_removed_between_devices():
    page = ("Art. 5º O órgão verificará a integridade.\n*CD261986821200* PRLP n.2\n"
            "Apresentação: 04/08/2026 18:06:05.563 - PLEN\nPRLP 2 => PL 5229/2025\n"
            "Assinado eletronicamente pelo(a) Dep. Fulano de Tal\n"
            "Para verificar a assinatura, acesse https://infoleg-autenticidade-assinatura.camara.leg.br/CD261986821200\n"
            "§ 1º A verificação será feita por servidor efetivo.")
    result = module._assemble([(3, None, page)], "test")
    assert "*CD" not in result.text and "Assinado" not in result.text and "=>" not in result.text
    assert result.text.startswith("Art. 5º") and "§ 1º A verificação" in result.text
    assert result.segments == [{"page": 3, "block": None, "start": 0, "end": len(result.text), "text": result.text}]
